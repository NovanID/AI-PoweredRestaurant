import openpyxl

wb = openpyxl.load_workbook('scratch/downloaded_timeline.xlsx')
ws_time = wb['Timeline']

for r in range(44, 80):
    no = ws_time.cell(r, 1).value
    kegiatan = ws_time.cell(r, 2).value
    marks = []
    for c in range(3, 26):
        cell = ws_time.cell(r, c)
        val = cell.value
        has_fill = False
        if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
            rgb = str(cell.fill.start_color.rgb)
            if rgb not in ['00000000', 'FFFFFFFF', 'None']:
                has_fill = True
        if val or has_fill:
            h = ws_time.cell(4, c).value or f"Col{c}"
            marks.append(h.replace('\n', ' '))
    if kegiatan:
        print(f"No {no} | {kegiatan[:50]}... | Weeks: {marks}")
