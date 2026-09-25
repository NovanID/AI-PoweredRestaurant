from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "Ridho-Magang-Timeline (1).xlsx"
OUTPUT = INPUT
SHEET = "internship-reports-2026-08-09.c"
TIMELINE = "Timeline"
SUMMARY = "Weekly Summary"

ACTIVITIES = [
    "Analisis kebutuhan dan perumusan konsep AI-Powered Restaurant SaaS",
    "Eksplorasi tools AI, agent, dan opsi model untuk chatbot restoran",
    "Perancangan arsitektur awal sistem AI Restaurant serta pemilihan technology stack",
    "Penyusunan roadmap pengembangan AI-Powered Restaurant berbasis SaaS",
    "Riset dan evaluasi model AI serta infrastruktur untuk kebutuhan chatbot restoran",
    "Sinkronisasi visi pengembangan asisten AI F&B dan penyesuaian roadmap multi-tenant",
    "Setup project structure dan initial development environment",
    "Pengembangan landing page utama dan komponen Navbar",
    "Pengembangan katalog menu dan manajemen data menu restoran",
    "Pengembangan form reservasi dan state management aplikasi",
    "Perencanaan scope, target MVP, dan roadmap AI-Powered Restaurant",
    "Pengembangan Admin Dashboard untuk manajemen reservasi dan meja",
    "Penyelesaian Admin Dashboard, POS, payment flow, dan initial commit",
    "Pembangunan fondasi AI: types, conversation state machine, dan context engine",
    "Penyelesaian AI orchestration, integrasi model, tool registry, dan Mock API",
    "Presentasi progres arsitektur AI, scope proyek, dan rencana implementasi",
    "Integrasi beberapa pilihan model AI melalui 9router",
    "Penerapan temporary table lock, anti-double booking, dan aturan jam operasional",
    "Pengembangan monitoring AI pada Admin Dashboard dan pengujian skenario reservasi",
    "Optimasi performa chat AI, widget chat, dan penanganan loading",
    "Presentasi progres minggu ke-4 dan evaluasi pengembangan AI chatbot",
    "Optimasi alur pemesanan takeaway melalui AI Chatbot hingga checkout",
    "Pembuatan tombol pembayaran cepat dan pop-up pembayaran online di AI Chat",
]


def main() -> None:
    workbook = load_workbook(INPUT)
    sheet = workbook[SHEET]

    timeline = workbook[TIMELINE]
    summary = workbook[SUMMARY]
    assert sheet["E1"].value == "Aktivitas"
    assert timeline["B4"].value == "Kegiatan"
    assert summary["D1"].value == "Aktivitas Utama"
    assert len(ACTIVITIES) == 23
    before = {
        (worksheet.title, cell.coordinate): cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    }
    target_styles = {
        **{(SHEET, f"E{row}"): sheet.cell(row=row, column=5).style_id for row in range(2, 25)},
        **{(TIMELINE, f"B{row}"): timeline.cell(row=row, column=2).style_id for row in range(5, 28)},
        **{(SUMMARY, f"D{row}"): summary.cell(row=row, column=4).style_id for row in range(2, 8)},
    }

    for row, activity in enumerate(ACTIVITIES, start=2):
        sheet.cell(row=row, column=5).value = activity
    for row, activity in enumerate(ACTIVITIES, start=5):
        timeline.cell(row=row, column=2).value = activity
    for row, activities in enumerate((ACTIVITIES[:1], ACTIVITIES[1:6], ACTIVITIES[6:11], ACTIVITIES[11:16], ACTIVITIES[16:21], ACTIVITIES[21:]), start=2):
        summary.cell(row=row, column=4).value = "; ".join(activities)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)

    saved = load_workbook(OUTPUT, data_only=False)
    saved_sheet = saved[SHEET]
    assert [saved_sheet.cell(row=row, column=5).value for row in range(2, 25)] == ACTIVITIES
    for (sheet_name, coordinate), style_id in target_styles.items():
        assert saved[sheet_name][coordinate].style_id == style_id
    for worksheet in saved.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if (worksheet.title == SHEET and cell.column == 5 and 2 <= cell.row <= 24) or (worksheet.title == TIMELINE and cell.column == 2 and 5 <= cell.row <= 27) or (worksheet.title == SUMMARY and cell.column == 4 and 2 <= cell.row <= 7):
                    continue
                assert cell.value == before[(worksheet.title, cell.coordinate)]


if __name__ == "__main__":
    main()
