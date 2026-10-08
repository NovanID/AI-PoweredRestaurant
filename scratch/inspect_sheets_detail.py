import openpyxl

wb = openpyxl.load_workbook('scratch/downloaded_timeline.xlsx', data_only=True)
print("=== SHEETS IN WORKBOOK ===")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"- {name} (max_row={ws.max_row}, max_column={ws.max_column})")

print("\n=== SHEET: Weekly Summary ===")
ws_sum = wb['Weekly Summary']
for r in range(1, min(15, ws_sum.max_row + 1)):
    row = [ws_sum.cell(r, c).value for c in range(1, min(10, ws_sum.max_column + 1))]
    print(f"R{r}: {row}")

print("\n=== SHEET: internship-reports-2026-08-09.c ===")
ws_rep = wb['internship-reports-2026-08-09.c']
for r in range(1, min(25, ws_rep.max_row + 1)):
    row = [ws_rep.cell(r, c).value for c in range(1, min(10, ws_rep.max_column + 1))]
    print(f"R{r}: {row}")
