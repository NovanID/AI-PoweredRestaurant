import openpyxl

wb = openpyxl.load_workbook('scratch/downloaded_timeline.xlsx', data_only=True)
ws_rep = wb['internship-reports-2026-08-09.c']
print(f"Max rows in internship-reports: {ws_rep.max_row}")
for r in range(23, ws_rep.max_row + 1):
    date_val = ws_rep.cell(r, 1).value
    status = ws_rep.cell(r, 4).value
    act = ws_rep.cell(r, 5).value
    if date_val or act:
        print(f"R{r} | Date: {date_val} | Status: {status} | Act: {act[:60] if act else 'None'}")
