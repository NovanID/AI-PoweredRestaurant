import openpyxl

wb = openpyxl.load_workbook('scratch/downloaded_timeline.xlsx')
ws_time = wb['Timeline']

print(f"Timeline max_row: {ws_time.max_row}, max_col: {ws_time.max_column}")
# Print header
headers = [ws_time.cell(4, c).value for c in range(1, 26)]
print("Headers:", headers)

for r in range(5, min(45, ws_time.max_row + 1)):
    no = ws_time.cell(r, 1).value
    kegiatan = ws_time.cell(r, 2).value
    # find marked columns
    marks = []
    for c in range(3, 26):
        cell = ws_time.cell(r, c)
        val = cell.value
        # check fill
        has_fill = False
        if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
            # check if not white or transparent
            rgb = str(cell.fill.start_color.rgb)
            if rgb not in ['00000000', 'FFFFFFFF', 'None']:
                has_fill = True
        if val or has_fill:
            h = ws_time.cell(4, c).value or f"Col{c}"
            marks.append(h.replace('\n', ' '))
    if kegiatan:
        print(f"No {no} | {kegiatan[:50]}... | Weeks: {marks}")
