import openpyxl
from datetime import datetime

wb = openpyxl.load_workbook('Ridho-MagangReports-GI (2).xlsx')
ws = wb.active

entries = [
    {
        'row': 18,
        'date': datetime(2026, 8, 31),
        'activity': 'Mencoba dan menghubungkan beberapa pilihan model AI alternatif melalui 9router',
        'progress': '1. Mencoba beberapa pilihan model AI lain (seperti Qwen dan DeepSeek) melalui 9router agar aplikasi tidak bergantung pada satu AI saja dan menghindari batasan kuota gratis Gemini.\n2. Mengubah kode koneksi AI di `lib/ai/gemini-client.ts` supaya bisa gonta-ganti model dan alamat API dengan mudah cukup lewat pengaturan `.env.local` / `.env.example`.\n3. Membuat file uji coba di `scratch/check-models.ts`, `scratch/find-active-model.ts`, dan `scratch/test-gemini-live.ts` untuk mengetes kecepatan respon dan ketepatan AI dalam menjawab pertanyaan.\n4. Menambahkan pembersih teks pada `lib/ai/gemini-client.ts` untuk membersihkan format jawaban dari 9router yang ada teks tambahannya agar tidak membuat aplikasi error.',
        'obstacle': 'Format jawaban dari 9router sempat ada teks tambahan yang membuat aplikasi error, dan beberapa model AI kuota gratisnya sempat habis.',
        'solution': 'Menambahkan pembersih teks otomatis di `lib/ai/gemini-client.ts` sebelum jawaban diproses oleh aplikasi, serta menyiapkan opsi model AI cadangan.',
        'next_plan': 'Membuat sistem reservasi yang lebih aman di `lib/domain/reservation-service.ts` agar tidak terjadi bentrok atau dobel booking meja.'
    },
    {
        'row': 19,
        'date': datetime(2026, 9, 1),
        'activity': 'Membuat sistem penguncian meja sementara (anti dobel booking) dan aturan jam buka restoran',
        'progress': '1. Membuat fitur penguncian meja sementara di `lib/domain/reservation-service.ts`. Saat pelanggan minta pesan meja lewat AI, meja tersebut dikunci sementara (5-10 menit) sampai pelanggan mengetik konfirmasi "Ya", sehingga tidak bisa diambil orang lain.\n2. Menambahkan aturan jam operasional restoran di `lib/domain/business-rules.ts` (buka jam 10.00 - 22.00) dan kapasitas meja, sehingga AI akan otomatis menolak jika ada yang memesan di luar jam buka.\n3. Menambahkan sistem pengaman di `lib/infrastructure/idempotency.ts` agar jika pelanggan tidak sengaja mengirim pesan berulang-ulang, pesanan tidak tercatat ganda.\n4. Menyiapkan sistem notifikasi event di `lib/infrastructure/event-bus.ts` untuk mencatat status meja yang sedang dikunci, disetujui, atau dibatalkan.',
        'obstacle': 'Takut terjadi bentrok meja (dobel booking) jika ada dua pelanggan yang memilih meja yang sama di waktu yang bersamaan.',
        'solution': 'Meja langsung dikunci sementara lewat `lib/domain/reservation-service.ts` saat ditawarkan ke pelanggan, dan baru resmi tersimpan ke database setelah pelanggan menyetujui ringkasan pesanannya.',
        'next_plan': 'Membuat halaman monitoring di dashboard admin (`components/AIChatMonitoringTab.tsx`) untuk melihat riwayat percakapan AI.'
    },
    {
        'row': 20,
        'date': datetime(2026, 9, 2),
        'activity': 'Membuat halaman monitoring AI di Admin Dashboard dan melakukan uji coba skenario reservasi',
        'progress': '1. Membuat tab baru di dashboard admin (`components/AIChatMonitoringTab.tsx`) dan menghubungkannya ke `components/AdminDashboard.tsx` supaya staf restoran bisa melihat langsung percakapan pelanggan dengan AI, status meja yang sedang dikunci, dan aksi apa saja yang dijalankan AI.\n2. Menambahkan modul pencatatan otomatis di `lib/infrastructure/observability.ts` untuk melihat berapa lama AI menjawab dan apakah ada pesan error.\n3. Melakukan tes pada file `scratch/test-orchestration.ts` untuk berbagai situasi: pelanggan minta menu pedas budget 40 ribu, pelanggan pesan meja, pelanggan konfirmasi, pelanggan cek nomor tiket, hingga pelanggan yang iseng memesan jam 3 subuh (otomatis ditolak).',
        'obstacle': 'Respon AI terasa agak lambat (menunggu di atas 3 detik) karena sistem harus memanggil AI sebanyak dua kali untuk satu pertanyaan.',
        'solution': 'Merencanakan perbaikan alur di `lib/ai/orchestrator.ts` agar AI bisa langsung memberikan jawaban teks sekaligus memproses data dalam satu kali jalan saja.',
        'next_plan': 'Mempercepat respon chat AI di `lib/ai/orchestrator.ts`, menghapus jeda loading di `components/AIChatWidget.tsx`, dan menambah batas waktu tunggu.'
    },
    {
        'row': 21,
        'date': datetime(2026, 9, 3),
        'activity': 'Mempercepat respon chat AI, merapikan tampilan widget chat, dan mengatasi masalah loading',
        'progress': '1. Mengubah cara kerja AI di `lib/ai/orchestrator.ts` dan `lib/ai/gemini-client.ts` agar langsung memberikan jawaban dan aksi dalam 1 langkah saja (tidak perlu 2 kali bolak-balik), sehingga waktu tunggu berkurang hingga separuhnya.\n2. Menghapus jeda loading buatan di tampilan chat `components/AIChatWidget.tsx` supaya pesan AI langsung muncul seketika di layar pengguna.\n3. Menambahkan batas waktu tunggu maksimal 15 detik di `lib/ai-assistant-service.ts` dan `lib/ai/gemini-client.ts`. Jika internet lambat atau server AI bermasalah, aplikasi tidak akan macet (nge-hang) dan menampilkan pesan ramah ke pengguna.\n4. Menyesuaikan konfigurasi di `package.json` (`next dev --webpack`) untuk mengatasi error tampilan CSS saat aplikasi dijalankan.',
        'obstacle': 'Khawatir chat macet atau loading terus-menerus tanpa henti jika server AI lambat merespons atau sedang error.',
        'solution': 'Memasang batas waktu (timeout) 15 detik di `lib/ai-assistant-service.ts`, jadi jika server AI macet, sistem otomatis membatalkan proses dan memberi tahu pengguna secara sopan.',
        'next_plan': 'Menyiapkan bahan dan demo dari file `AI-architecture.md` dan `scratch/test-orchestration.ts` untuk presentasi evaluasi mingguan dengan mentor.'
    },
    {
        'row': 22,
        'date': datetime(2026, 9, 4),
        'activity': 'Evaluasi mingguan (Presentasi Minggu ke-4) dan demo hasil pengembangan fitur AI restoran',
        'progress': '1. Menyiapkan skenario demo dari `scratch/test-orchestration.ts` dan `app/page.tsx` untuk mentor: menunjukkan percakapan pelanggan dengan AI, pemesanan meja otomatis yang anti bentrok, rekomendasi menu, dan cara staf melihat monitoring percakapan di dashboard admin.\n2. Membahas kemajuan sistem AI berdasarkan dokumen `AI-architecture.md` dan `Insight-project.md` yang sekarang lebih cepat dan stabil dibanding minggu sebelumnya.\n3. Merapikan catatan dan dokumentasi laporan kemajuan di file `Ridho-MagangReports-GI (2).xlsx`.',
        'obstacle': 'Perlu memastikan koneksi dan server AI tetap lancar saat sesi demo di depan mentor.',
        'solution': 'Menyiapkan data simulasi di `lib/mock-data.ts` dan mencoba semua alur percakapan terlebih dahulu sebelum sesi pertemuan dimulai.',
        'next_plan': 'Melanjutkan pengembangan tahap berikutnya: menghubungkan chat AI ke WhatsApp agar pelanggan bisa langsung chat dari WA.'
    }
]

from openpyxl.styles import Font, Alignment

font_roboto = Font(name='Roboto', size=10)
align_wrap = Alignment(wrap_text=True, vertical='top')
align_date = Alignment(wrap_text=True, vertical='top', horizontal='center')

for entry in entries:
    r = entry['row']
    c1 = ws.cell(row=r, column=1, value=entry['date'])
    c1.number_format = 'yyyy-mm-dd'
    c1.font = font_roboto
    c1.alignment = align_date
    
    c2 = ws.cell(row=r, column=2, value=entry['activity'])
    c2.font = font_roboto
    c2.alignment = align_wrap
    
    c3 = ws.cell(row=r, column=3, value=entry['progress'])
    c3.font = font_roboto
    c3.alignment = align_wrap
    
    c4 = ws.cell(row=r, column=4, value=entry['obstacle'])
    c4.font = font_roboto
    c4.alignment = align_wrap
    
    c5 = ws.cell(row=r, column=5, value=entry['solution'])
    c5.font = font_roboto
    c5.alignment = align_wrap
    
    c6 = ws.cell(row=r, column=6, value=entry['next_plan'])
    c6.font = font_roboto
    c6.alignment = align_wrap

# Bersihkan sisa baris 23 dan 24
for r in [23, 24]:
    for c in range(1, 7):
        ws.cell(row=r, column=c, value=None)

wb.save('Ridho-MagangReports-GI (2).xlsx')
print('Berhasil update Excel dengan nama file lengkap!')
