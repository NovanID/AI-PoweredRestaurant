import urllib.request
import openpyxl

url = 'https://docs.google.com/spreadsheets/d/1rPqAjEaHe2EeTQQArVUIR8JEj2khGngJE8p2s088-Ak/export?format=xlsx'
req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
content = urllib.request.urlopen(req).read()
with open('scratch/downloaded_timeline.xlsx', 'wb') as f:
    f.write(content)

wb = openpyxl.load_workbook('scratch/downloaded_timeline.xlsx')
print('Sheet names:', wb.sheetnames)
ws = wb.active
for r in range(1, 40):
    row_vals = [ws.cell(r, c).value for c in range(1, 26)]
    # check if any cell has fill
    fills = [ws.cell(r, c).fill.start_color.rgb for c in range(1, 26) if ws.cell(r, c).fill and ws.cell(r, c).fill.start_color and ws.cell(r, c).fill.start_color.rgb]
    if any(row_vals):
        print(f"Row {r}: {row_vals[:3]} | Fills: {len(fills)}")
